#!/usr/bin/env python3
"""Build the citation-verification spreadsheet and a markdown summary.

Usage:
  make_report.py verify_all.json -o citation_verification.xlsx [--actions actions.json]

verify_all.json: merged output of the verification agents (list of records, see
references/verification-agent-prompt.md for the schema).
actions.json (optional): {"<key>": "what was changed in the manuscript"}.

Requires: openpyxl.
"""
import argparse, json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

FILLS = {'NOT SUPPORTED': 'F4CCCC', 'PARTLY SUPPORTED': 'FFF2CC', 'UNVERIFIABLE': 'E7E6E6'}


def verdict_of(c):
    return (c.get('verdict') or '').upper().split(' (')[0].split(':')[0].strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('verify'); ap.add_argument('-o', '--out', default='citation_verification.xlsx')
    ap.add_argument('--actions'); ap.add_argument('--note', default='')
    a = ap.parse_args()
    recs = json.load(open(a.verify))
    actions = json.load(open(a.actions)) if a.actions else {}

    wb = Workbook(); ws = wb.active; ws.title = 'Claims'
    ws.append(['Key', 'Reference (published record)', 'DOI', 'Version checked', 'PDF saved',
               'Claim we make', 'Verdict', 'Evidence', 'Notes', 'Action taken'])
    wrap = Alignment(wrap_text=True, vertical='top')
    for e in recs:
        au = e.get('authors')
        if isinstance(au, list):
            au = ', '.join(f"{x.get('family','')}, {x.get('given','')}" for x in au)
        ref = f"{au} ({e.get('year')}). {e.get('title')}. {e.get('journal') or e.get('container') or e.get('publisher') or ''}"
        if e.get('volume'): ref += f", {e['volume']}"
        if e.get('issue'): ref += f"({e['issue']})"
        if e.get('pages'): ref += f": {e['pages']}"
        for c in e.get('claims', []):
            ws.append([e['key'], ref, e.get('doi') or '', e.get('version_checked') or '',
                       'yes' if e.get('pdf_downloaded') else 'no', c.get('claim'), c.get('verdict'),
                       c.get('evidence'), c.get('notes'), actions.get(e['key'], 'No change.')])
    for i, w in enumerate([18, 45, 22, 30, 9, 50, 16, 50, 50, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9D9D9')
    for row in ws.iter_rows(min_row=2):
        v = (row[6].value or '').upper()
        for k, col in FILLS.items():
            if v.startswith(k):
                row[6].fill = PatternFill('solid', fgColor=col)
    ws.freeze_panes = 'A2'

    cnt = Counter(verdict_of(c) for e in recs for c in e.get('claims', []))
    s = wb.create_sheet('Summary')
    s.append(['Verdict', 'Claims'])
    for k, v in sorted(cnt.items()):
        s.append([k, v])
    s.append([])
    vc = lambda pref: sum(1 for e in recs if str(e.get('version_checked', '')).lower().startswith(pref))
    s.append(['References checked', len(recs)])
    s.append(['PDFs saved', sum(1 for e in recs if e.get('pdf_downloaded'))])
    s.append(['Checked against published version', vc('published')])
    s.append(['Checked against preprint / working paper', vc('preprint')])
    s.append(['Abstract, book description, or secondary source only', vc('abstract') + vc('book') + vc('secondary')])
    if a.note:
        s.append([]); s.append(['Note', a.note])
    s.column_dimensions['A'].width = 45; s.column_dimensions['B'].width = 90
    for row in s.iter_rows():
        for c in row:
            c.alignment = wrap
    for c in s[1]:
        c.font = Font(bold=True)
    wb.save(a.out)

    # markdown summary to stdout
    print(f'## Citation check: {len(recs)} references, {sum(cnt.values())} claims')
    print(', '.join(f'{k.lower()}: {v}' for k, v in sorted(cnt.items())))
    for label in ('NOT SUPPORTED', 'PARTLY SUPPORTED', 'UNVERIFIABLE'):
        items = [(e['key'], c) for e in recs for c in e.get('claims', []) if verdict_of(c) == label]
        if items:
            print(f'\n### {label.title()}')
            for k, c in items:
                print(f"- **{k}** — {c.get('claim')}\n  - {c.get('notes') or c.get('evidence')}")
    print(f'\nSpreadsheet: {a.out}')


if __name__ == '__main__':
    main()
